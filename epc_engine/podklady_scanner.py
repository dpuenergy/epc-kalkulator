"""
Scanner složky s podklady projektu.

Prochází PDF, Word a Excel soubory, extrahuje z nich text a připravuje
kontext pro AI generování popisů v pasportu.

Skenovaná PDF jsou automaticky detekována a nabídnuta ke zpracování
přes Claude vision API.
"""
from __future__ import annotations

import base64
import re
from pathlib import Path

SUPPORTED_EXTS = {".pdf", ".docx", ".doc", ".xlsx", ".xls"}
MAX_CHARS_PER_FILE = 6000   # ~1500 tokenů
MAX_PAGES_PDF = 12          # max stránek z jednoho PDF
MAX_SHEETS_XLSX = 3         # max listů z jednoho XLSX
SCAN_THRESHOLD = 80         # znaků/stránku – méně = sken


# ─────────────────────────────────────────────────────────────────────────────
# Skenování složky
# ─────────────────────────────────────────────────────────────────────────────

def scan_folder(folder: str | Path) -> list[dict]:
    """
    Rekurzivně projde složku a vrátí seznam podporovaných souborů.

    Každý záznam: {path, name, rel_path, ext, size_kb}
    Seřazeno: nejprve PDF, pak Word, pak Excel; abecedně v rámci skupiny.
    """
    folder = Path(folder)
    if not folder.is_dir():
        return []

    results = []
    for p in folder.rglob("*"):
        if not p.is_file():
            continue
        ext = p.suffix.lower()
        if ext not in SUPPORTED_EXTS:
            continue
        # přeskočit dočasné soubory Office (~$...)
        if p.name.startswith("~$"):
            continue
        results.append({
            "path": str(p),
            "name": p.name,
            "rel_path": str(p.relative_to(folder)),
            "ext": ext,
            "size_kb": round(p.stat().st_size / 1024),
        })

    _order = {".pdf": 0, ".docx": 1, ".doc": 1, ".xlsx": 2, ".xls": 2}
    results.sort(key=lambda r: (_order.get(r["ext"], 9), r["rel_path"].lower()))
    return results


# ─────────────────────────────────────────────────────────────────────────────
# Extrakce textu
# ─────────────────────────────────────────────────────────────────────────────

def _extract_pdf_text(path: str) -> tuple[str, bool]:
    """
    Extrahuje text z PDF pomocí pdfplumber.
    Vrátí (text, je_sken).
    je_sken=True pokud průměrný počet znaků/stránku < SCAN_THRESHOLD.
    """
    try:
        import pdfplumber
    except ImportError:
        return "", False

    parts: list[str] = []
    page_count = 0
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages[:MAX_PAGES_PDF]:
            t = (page.extract_text() or "").strip()
            parts.append(t)
            page_count += 1

    total_chars = sum(len(t) for t in parts)
    is_scan = page_count > 0 and (total_chars / page_count) < SCAN_THRESHOLD
    text = "\n\n".join(t for t in parts if t)
    return text[:MAX_CHARS_PER_FILE], is_scan


def _extract_pdf_scan_claude(path: str, api_key: str) -> str:
    """
    Extrahuje text ze skenovaného PDF pomocí Claude vision (PyMuPDF → obrázky).
    """
    try:
        import fitz  # PyMuPDF
        import anthropic
    except ImportError as exc:
        raise RuntimeError(f"Chybí balíček: {exc}") from exc

    doc = fitz.open(path)
    client = anthropic.Anthropic(api_key=api_key)
    parts: list[str] = []

    for i in range(min(len(doc), MAX_PAGES_PDF)):
        page = doc[i]
        pix = page.get_pixmap(dpi=150)
        img_b64 = base64.standard_b64encode(pix.tobytes("png")).decode()

        msg = client.messages.create(
            model="claude-haiku-4-5-20251001",
            max_tokens=800,
            messages=[{
                "role": "user",
                "content": [
                    {
                        "type": "image",
                        "source": {
                            "type": "base64",
                            "media_type": "image/png",
                            "data": img_b64,
                        },
                    },
                    {
                        "type": "text",
                        "text": (
                            "Extrahuj veškerý text z této stránky dokumentu. "
                            "Zachovej strukturu a pořadí textu. "
                            "Pokud jde o tabulku, zachovej ji jako text oddělený svislítkem."
                        ),
                    },
                ],
            }],
        )
        parts.append(msg.content[0].text.strip())

    doc.close()
    return "\n\n".join(parts)[:MAX_CHARS_PER_FILE * 2]


def _extract_docx(path: str) -> str:
    try:
        from docx import Document
    except ImportError:
        return ""
    doc = Document(path)
    lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    return "\n".join(lines)[:MAX_CHARS_PER_FILE]


def _extract_xlsx(path: str) -> str:
    try:
        import openpyxl
    except ImportError:
        return ""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in list(wb.worksheets)[:MAX_SHEETS_XLSX]:
        rows: list[str] = []
        for row in sheet.iter_rows(values_only=True):
            vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
            if vals:
                rows.append(" | ".join(vals))
            if sum(len(r) for r in rows) > MAX_CHARS_PER_FILE:
                break
        if rows:
            parts.append(f"[List: {sheet.title}]\n" + "\n".join(rows))
    wb.close()
    return "\n\n".join(parts)[:MAX_CHARS_PER_FILE]


def extract_text(file_info: dict) -> tuple[str, bool]:
    """
    Extrahuje text ze souboru popsaného záznamem ze scan_folder().
    Vrátí (text, je_sken_pdf).
    """
    path = file_info["path"]
    ext = file_info["ext"]
    if ext == ".pdf":
        return _extract_pdf_text(path)
    if ext in (".docx", ".doc"):
        return _extract_docx(path), False
    if ext in (".xlsx", ".xls"):
        return _extract_xlsx(path), False
    return "", False


def extract_text_scan(file_info: dict, api_key: str) -> str:
    """Extrahuje text ze skenovaného PDF přes Claude vision."""
    return _extract_pdf_scan_claude(file_info["path"], api_key)


# ─────────────────────────────────────────────────────────────────────────────
# Formátování kontextu pro AI prompt
# ─────────────────────────────────────────────────────────────────────────────

def _clean(text: str) -> str:
    """Odstraní nadbytečné prázdné řádky."""
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def sestavit_kontext(extracts: dict[str, str]) -> str:
    """
    Ze slovníku {název_souboru: text} sestaví kontextový blok pro AI prompt.
    """
    if not extracts:
        return ""
    parts = []
    for name, text in extracts.items():
        if text.strip():
            parts.append(f"--- {name} ---\n{_clean(text)}")
    if not parts:
        return ""
    return (
        "Níže jsou výtahy z dostupných podkladů k objektu. "
        "Použij je jako doplňující zdroj informací při psaní popisu:\n\n"
        + "\n\n".join(parts)
    )
